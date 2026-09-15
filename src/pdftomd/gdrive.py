"""Google Drive support for `.gdoc` / `.gsheet` / `.gslides` stub files.

Google Drive for Desktop stores native Google documents as tiny JSON stubs::

    {"doc_id": "1oCH83…", "resource_key": "", "email": "…"}

The content lives online only, so this module fetches it through the Drive
API (read-only scope) and turns it into Markdown:

* Docs    -> exported natively as ``text/markdown``
* Sheets  -> exported as ``.xlsx`` and rendered as one GFM table per tab
* Slides  -> exported as PDF, then handed to the selected conversion engine

Change detection uses the remote ``version``/``modifiedTime`` as fingerprint,
since the stub bytes never change when the document is edited.
"""

from __future__ import annotations

import io
import json
import os
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Protocol

from .converters.base import ConversionError

SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]
API = "https://www.googleapis.com/drive/v3/files"
STUB_EXTENSIONS = frozenset({".gdoc", ".gsheet", ".gslides"})

MIME_DOC = "application/vnd.google-apps.document"
MIME_SHEET = "application/vnd.google-apps.spreadsheet"
MIME_SLIDES = "application/vnd.google-apps.presentation"
MIME_SHORTCUT = "application/vnd.google-apps.shortcut"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MIME_PDF = "application/pdf"
MIME_MD = "text/markdown"

ENGINE_NAME = "gdrive"  # recorded in the manifest when no conversion engine was involved


class GoogleDriveError(ConversionError):
    """Auth, network or API failure while talking to Google Drive."""


# -- configuration / auth ------------------------------------------------------


def config_dir() -> Path:
    return Path(os.environ.get("PDFTOMD_CONFIG_DIR", Path.home() / ".config" / "pdftomd"))


def client_secret_path() -> Path:
    return config_dir() / "client_secret.json"


def token_path() -> Path:
    return config_dir() / "google-token.json"


def has_token() -> bool:
    return token_path().exists()


def login(client_secret: Path | None = None) -> Path:
    """Run the OAuth desktop flow once and cache a refreshable token. Returns the token path."""
    try:
        from google_auth_oauthlib.flow import InstalledAppFlow
    except ImportError as exc:  # pragma: no cover - depends on env
        raise GoogleDriveError("Install with `pip install 'pdftomd[gdrive]'`.") from exc
    secret = Path(client_secret) if client_secret else client_secret_path()
    if not secret.exists():
        raise GoogleDriveError(
            f"OAuth client secret not found at {secret}. Create a 'Desktop app' OAuth client in the "
            "Google Cloud console (Drive API enabled) and download it there."
        )
    flow = InstalledAppFlow.from_client_secrets_file(str(secret), SCOPES)
    creds = flow.run_local_server(port=0, open_browser=True)
    out = token_path()
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(creds.to_json(), "utf-8")
    os.chmod(out, 0o600)
    return out


def load_credentials():
    try:
        from google.auth.transport.requests import Request
        from google.oauth2.credentials import Credentials
    except ImportError as exc:  # pragma: no cover - depends on env
        raise GoogleDriveError("Install with `pip install 'pdftomd[gdrive]'`.") from exc
    path = token_path()
    if not path.exists():
        raise GoogleDriveError(f"Not logged in to Google Drive (no token at {path}). Run `pdftomd gdrive login`.")
    creds = Credentials.from_authorized_user_file(str(path), SCOPES)
    if not creds.valid:
        if not (creds.expired and creds.refresh_token):
            raise GoogleDriveError("Google Drive token is invalid. Run `pdftomd gdrive login` again.")
        try:
            creds.refresh(Request())
        except Exception as exc:
            raise GoogleDriveError(f"Could not refresh Google Drive token: {exc}. Run `pdftomd gdrive login`.") from exc
        path.write_text(creds.to_json(), "utf-8")
    return creds


# -- stubs ---------------------------------------------------------------------


@dataclass(frozen=True)
class Stub:
    doc_id: str
    resource_key: str = ""

    @property
    def headers(self) -> dict[str, str]:
        return {"X-Goog-Drive-Resource-Keys": f"{self.doc_id}/{self.resource_key}"} if self.resource_key else {}


def is_stub(path: str | Path) -> bool:
    return Path(path).suffix.lower() in STUB_EXTENSIONS


def read_stub(path: str | Path) -> Stub:
    path = Path(path)
    try:
        data = json.loads(path.read_text("utf-8"))
    except (OSError, ValueError) as exc:
        raise GoogleDriveError(f"{path.name}: not a valid Google Drive stub: {exc}") from exc
    doc_id = data.get("doc_id")
    if not doc_id and (url := data.get("url")):  # older Drive clients stored an URL instead
        m = re.search(r"/d/([A-Za-z0-9_-]+)", url)
        doc_id = m.group(1) if m else None
    if not doc_id:
        raise GoogleDriveError(f"{path.name}: no doc_id in Google Drive stub")
    return Stub(doc_id=doc_id, resource_key=data.get("resource_key") or "")


# -- API client ----------------------------------------------------------------


class DriveClient(Protocol):
    def metadata(self, stub: Stub) -> dict[str, Any]: ...
    def export(self, stub: Stub, mime: str) -> bytes: ...


class HttpDriveClient:
    """Thin wrapper over the Drive v3 REST API using an authorized requests session."""

    FIELDS = "id,name,mimeType,modifiedTime,version,shortcutDetails"

    def __init__(self, credentials=None) -> None:
        self._creds = credentials
        self._session = None

    def _get_session(self):
        if self._session is None:
            from google.auth.transport.requests import AuthorizedSession

            self._session = AuthorizedSession(self._creds or load_credentials())
        return self._session

    def _get(self, url: str, stub: Stub, **params):
        session = self._get_session()
        try:
            resp = session.get(url, params={"supportsAllDrives": "true", **params}, headers=stub.headers, timeout=120)
        except Exception as exc:
            raise GoogleDriveError(f"Google Drive request failed: {exc}") from exc
        if resp.status_code != 200:
            raise GoogleDriveError(f"Google Drive API {resp.status_code} for {stub.doc_id}: {resp.text[:300]}")
        return resp

    def metadata(self, stub: Stub) -> dict[str, Any]:
        meta = self._get(f"{API}/{stub.doc_id}", stub, fields=self.FIELDS).json()
        if meta.get("mimeType") == MIME_SHORTCUT and (target := meta.get("shortcutDetails", {}).get("targetId")):
            meta = self._get(f"{API}/{target}", stub, fields=self.FIELDS).json()
        return meta

    def export(self, stub: Stub, mime: str) -> bytes:
        return self._get(f"{API}/{stub.doc_id}/export", stub, mimeType=mime).content


# -- resolving -----------------------------------------------------------------


@dataclass(frozen=True)
class RemoteInfo:
    doc_id: str
    name: str
    mime: str
    fingerprint: str
    needs_engine: bool  # True when the export must go through a conversion engine (Slides -> PDF)


@dataclass(frozen=True)
class Document:
    """Result of resolving a stub: either final Markdown or bytes for an engine."""

    filename: str
    fingerprint: str
    markdown: str | None = None
    data: bytes | None = None


class GoogleDriveResolver:
    extensions = STUB_EXTENSIONS

    def __init__(self, client: DriveClient | None = None) -> None:
        self.client: DriveClient = client or HttpDriveClient()
        self._meta_cache: dict[str, dict[str, Any]] = {}

    def clear_cache(self) -> None:
        self._meta_cache.clear()

    def _metadata(self, stub: Stub) -> dict[str, Any]:
        if stub.doc_id not in self._meta_cache:
            self._meta_cache[stub.doc_id] = self.client.metadata(stub)
        return self._meta_cache[stub.doc_id]

    def describe(self, path: str | Path) -> RemoteInfo:
        """Cheap metadata lookup used for change detection (one API call, cached)."""
        stub = read_stub(path)
        meta = self._metadata(stub)
        mime = meta.get("mimeType", "")
        if mime not in (MIME_DOC, MIME_SHEET, MIME_SLIDES):
            raise GoogleDriveError(f"{Path(path).name}: unsupported Google Drive type {mime!r}")
        fingerprint = f"v{meta.get('version', '?')}@{meta.get('modifiedTime', '?')}"
        return RemoteInfo(meta["id"], meta.get("name", Path(path).stem), mime, fingerprint, needs_engine=mime == MIME_SLIDES)

    def resolve(self, path: str | Path) -> Document:
        """Fetch the document content. Returns Markdown for Docs/Sheets, PDF bytes for Slides."""
        stub = read_stub(path)
        info = self.describe(path)
        target = Stub(info.doc_id, stub.resource_key)
        stem = Path(path).stem
        if info.mime == MIME_DOC:
            try:
                md = self.client.export(target, MIME_MD).decode("utf-8")
                return Document(f"{stem}.md", info.fingerprint, markdown=_tidy(md))
            except GoogleDriveError as exc:
                if "exportSizeLimitExceeded" not in str(exc):
                    raise
                # Too large for Markdown export: fall back to PDF through the engine.
                return Document(f"{stem}.pdf", info.fingerprint, data=self.client.export(target, MIME_PDF))
        if info.mime == MIME_SHEET:
            xlsx = self.client.export(target, MIME_XLSX)
            return Document(f"{stem}.md", info.fingerprint, markdown=xlsx_to_markdown(xlsx, title=info.name))
        return Document(f"{stem}.pdf", info.fingerprint, data=self.client.export(target, MIME_PDF))


# -- rendering -----------------------------------------------------------------


def _tidy(md: str) -> str:
    md = md.replace("\r\n", "\n")
    md = re.sub(r"\n{3,}", "\n\n", md)
    return md.strip() + "\n"


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).replace("|", "\\|").replace("\r\n", "\n").replace("\n", "<br>")
    return text.strip()


def xlsx_to_markdown(data: bytes, *, title: str | None = None, include_hidden: bool = False) -> str:
    """Render every worksheet of an .xlsx file as a GFM table (first row = header)."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on env
        raise GoogleDriveError("openpyxl is required for Sheets. Install with `pip install 'pdftomd[gdrive]'`.") from exc
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    if title:
        parts.append(f"# {title}\n")
    for ws in wb.worksheets:
        if ws.sheet_state != "visible" and not include_hidden:
            continue
        rows = [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(r)]
        parts.append(f"## {ws.title}\n")
        if not rows:
            parts.append("_(empty sheet)_\n")
            continue
        width = max(len(r) for r in rows)
        while width > 0 and not any(r[width - 1] if len(r) >= width else "" for r in rows):
            width -= 1
        rows = [r[:width] + [""] * (width - len(r)) for r in rows]
        header, body = rows[0], rows[1:]
        lines = ["| " + " | ".join(header) + " |", "|" + "---|" * width]
        lines += ["| " + " | ".join(r) + " |" for r in body]
        parts.append("\n".join(lines) + "\n")
    return "\n".join(parts)
